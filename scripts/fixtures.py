"""Generate original synthetic import fixtures. No personal or third-party designs."""
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

root=Path(__file__).resolve().parents[1]/'fixtures';root.mkdir(exist_ok=True)
svg='''<svg xmlns="http://www.w3.org/2000/svg" width="800" height="600" viewBox="0 0 800 600">
<defs><linearGradient id="blue"><stop stop-color="#2563eb"/><stop offset="1" stop-color="#60a5fa"/></linearGradient><clipPath id="clip"><rect x="0" y="0" width="180" height="100" rx="20"/></clipPath></defs>
<rect width="800" height="600" fill="#eaf0ff"/><rect x="40" y="40" width="720" height="520" rx="24" fill="white"/>
<text x="80" y="130" font-family="Arial" font-size="16" fill="#2563eb">ONE WORKSPACE STUDIO</text>
<text x="80" y="230" font-family="Georgia" font-size="64" fill="#0b2a63">Make it yours.</text>
<text x="80" y="290" font-family="Arial" font-size="24" fill="#5e6c84"><tspan>Editable text, </tspan><tspan font-weight="bold">real possibilities.</tspan></text>
<g id="nested" transform="translate(80 350)"><g transform="rotate(8 90 50)"><rect width="180" height="100" fill="url(#blue)"/><circle cx="150" cy="50" r="38" fill="#a9f1da"/></g></g>
<path d="M340 470 L400 350 L460 470 Z" fill="#2563eb" opacity="0.8"/>
</svg>'''
(root/'editable.svg').write_text(svg,encoding='utf8')
(root/'missing-font.svg').write_text(svg.replace('Georgia','Definitely Missing Typeface'),encoding='utf8')
(root/'outlined.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="150"><path d="M20 130 L60 20 L100 130 M40 90 L80 90" fill="none" stroke="#2563eb" stroke-width="12"/></svg>')
(root/'hostile.svg').write_text('''<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100" onload="window.__pwned=true"><script>window.__pwned=true</script><foreignObject><div>bad</div></foreignObject><image href="https://example.invalid/tracker"/><text x="10" y="50" font-size="20">Safe content</text></svg>''')
(root/'malformed.svg').write_text('<svg><g></svg>')
im=Image.new('RGB',(800,600),'white');d=ImageDraw.Draw(im);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48);d.text((70,90),'A fresh perspective',fill='#0b2a63',font=font);d.rectangle((70,220,350,380),fill='#2563eb');d.ellipse((450,230,620,400),fill='#60a5fa');d.text((70,480),'Made for your next idea',fill='#18233b',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',26))
for ext in ('png','jpg','webp'):im.save(root/('flattened.'+ext))
c=Canvas(str(root/'vector-text-scan.pdf'),pagesize=(600,450));c.setFillColorRGB(1,1,1);c.rect(0,0,600,450,fill=1,stroke=0);c.setFillColorRGB(.145,.388,.922);c.rect(50,90,180,120,fill=1,stroke=0);c.setFillColorRGB(.043,.165,.388);c.setFont('Helvetica',32);c.drawString(50,350,'Editable PDF text');c.setFont('Helvetica',16);c.drawString(50,310,'Vectors and selectable text');c.drawImage(ImageReader(im),(300),(60),width=240,height=180);c.showPage();c.drawImage(ImageReader(im),0,0,width=600,height=450);c.showPage();c.save()
pres=Presentation();pres.slide_width=Inches(10);pres.slide_height=Inches(7.5)
slide=pres.slides.add_slide(pres.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(.8),Inches(.8),Inches(8),Inches(1.3));p=box.text_frame.paragraphs[0];run=p.add_run();run.text='Your next idea';run.font.size=Pt(44);run.font.name='Georgia';run.font.color.rgb=RGBColor(11,42,99)
box=slide.shapes.add_textbox(Inches(.8),Inches(2),Inches(8),Inches(.6));p=box.text_frame.paragraphs[0];r=p.add_run();r.text='Real text. ';r.font.size=Pt(22);r=p.add_run();r.text='Editable layers.';r.font.bold=True;r.font.size=Pt(22)
group=slide.shapes.add_group_shape();sh=group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(1),Inches(3),Inches(3),Inches(2));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(37,99,235);sh.line.fill.background();sh=group.shapes.add_shape(MSO_SHAPE.OVAL,Inches(3),Inches(3.5),Inches(1),Inches(1));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(169,241,218);group.rotation=7
slide.shapes.add_picture(str(root/'flattened.png'),Inches(5),Inches(3),width=Inches(4))
slide2=pres.slides.add_slide(pres.slide_layouts[6]);table=slide2.shapes.add_table(2,2,Inches(1),Inches(1),Inches(8),Inches(3)).table;table.cell(0,0).text='First';table.cell(0,1).text='Second';table.cell(1,0).text='Editable';table.cell(1,1).text='Table cells';pres.save(root/'groups-table.pptx')
(root/'LICENSE.txt').write_text('Original synthetic fixtures by One Workspace Studio. CC0-1.0. No external assets.\n')
print('Generated fixtures in',root)
